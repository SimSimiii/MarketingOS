import asyncio
import subprocess
import sys

import pytest
from docx import Document
from pypdf import PdfWriter

from app.ingestion import extraction
from app.ingestion.exceptions import LoaderError


def test_parser_memory_is_bounded_in_a_disposable_process():
    code = """
from app.ingestion.process_limits import limit_memory
limit_memory(64 * 1024 * 1024)
try:
    data = bytearray(128 * 1024 * 1024)
except MemoryError:
    print('bounded')
else:
    raise SystemExit('Memory limit was not enforced')
"""
    result = subprocess.run([sys.executable, "-c", code], capture_output=True,
                            text=True, check=True, timeout=15)
    assert result.stdout.strip() == "bounded"


@pytest.mark.asyncio
async def test_docx_preserves_table_prices_in_document_order(tmp_path):
    path = tmp_path / "pricing.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Starter"
    table.cell(0, 1).text = "$29/month"
    document.add_paragraph("After table")
    document.save(path)
    result = await extraction.extract_document("docx", str(path))
    assert result["content"].index("Before") < result["content"].index("$29/month")
    assert result["content"].index("$29/month") < result["content"].index("After")


@pytest.mark.asyncio
async def test_pdf_page_limit_is_enforced_in_child_process(tmp_path):
    path = tmp_path / "pages.pdf"
    writer = PdfWriter()
    for _ in range(201):
        writer.add_blank_page(width=10, height=10)
    writer.write(path)
    with pytest.raises(LoaderError, match="200 page"):
        await extraction.extract_document("pdf", str(path))


@pytest.mark.asyncio
async def test_cancelled_parser_is_killed_and_releases_capacity(monkeypatch):
    started = asyncio.Event()
    class Process:
        returncode = None
        killed = False
        async def communicate(self):
            started.set()
            await asyncio.Future()
        def kill(self):
            self.killed = True
            self.returncode = -1
        async def wait(self):
            return self.returncode
    process = Process()
    async def spawn(*args, **kwargs):
        return process
    monkeypatch.setattr(extraction.asyncio, "create_subprocess_exec", spawn)
    task = asyncio.create_task(extraction.extract_document("pdf", "unused"))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert process.killed
    assert extraction._workers.acquire(blocking=False)
    assert extraction._workers.acquire(blocking=False)
    extraction._workers.release()
    extraction._workers.release()
